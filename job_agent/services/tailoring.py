from hashlib import sha256
from pathlib import Path
from uuid import uuid4

from docx import Document

from job_agent.config import settings
from job_agent.models import Application, CandidateProfile, Resume


def generate_ats_docx(
    profile: CandidateProfile,
    source_resume: Resume,
    application: Application,
) -> tuple[str, Path, str, list[str], str]:
    """Reformat source text and surface skills evidenced in resume and target job only."""
    job = application.job
    resume_text_lower = source_resume.text.lower()
    job_text = f"{job.title} {job.description} {' '.join(job.required_skills or [])}".lower()
    selected = [
        skill for skill in (profile.skills or [])
        if skill.strip() and skill.lower() in resume_text_lower and skill.lower() in job_text
    ]

    version_id = str(uuid4())
    filename = f"resume-{application.id}-{version_id[:8]}.docx"
    output_dir = (settings.data_dir / "resume_versions").resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    path = output_dir / filename

    document = Document()
    if selected:
        document.add_heading("Relevant Skills", level=1)
        document.add_paragraph(", ".join(selected))
    document.add_heading("Resume", level=1)
    for line in source_resume.text.splitlines():
        if line.strip():
            document.add_paragraph(line.strip())
    document.save(path)

    content_hash = sha256(path.read_bytes()).hexdigest()
    summary = (
        "Generated ATS-friendly DOCX from the immutable uploaded resume. "
        "Source wording was not rewritten; relevant skills are selected only when present "
        "in both the source resume and the job description."
    )
    return version_id, path, content_hash, selected, summary
